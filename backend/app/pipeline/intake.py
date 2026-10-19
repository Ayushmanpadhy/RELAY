import io
import json
import logging
import os
import re
from pathlib import Path
from typing import Literal

import docx
import openpyxl
import pypdf
from fastapi import UploadFile
from PIL import Image

from app.config import settings
from app.models import FileManifest, Manifest

logger = logging.getLogger("relay.pipeline.intake")

ALLOWED_EXTENSIONS = {
    ".txt",
    ".md",
    ".csv",
    ".json",
    ".pdf",
    ".docx",
    ".xlsx",
    ".png",
    ".jpg",
    ".jpeg",
    ".webp",
}
MAX_FILE_BYTES = 25 * 1024 * 1024  # 25 MB
MAX_TOTAL_BYTES = 50 * 1024 * 1024  # 50 MB
MAX_FILES_COUNT = 10


class IntakeError(Exception):
    pass


async def process_intake(job_id: str, files: list[UploadFile]) -> tuple[Manifest, list[str]]:
    if len(files) > MAX_FILES_COUNT:
        raise IntakeError(f"Maximum {MAX_FILES_COUNT} files allowed. Received {len(files)}.")

    job_upload_dir = settings.UPLOAD_DIR / job_id
    job_upload_dir.mkdir(parents=True, exist_ok=True)
    raw_dir = job_upload_dir / "raw"
    raw_dir.mkdir(parents=True, exist_ok=True)
    extracted_dir = job_upload_dir / "extracted"
    extracted_dir.mkdir(parents=True, exist_ok=True)

    manifest_files: list[FileManifest] = []
    saved_file_paths: list[str] = []
    total_bytes = 0
    all_extracted_lines: list[str] = []
    used_names: set[str] = set()

    for idx, file in enumerate(files):
        if not file.filename:
            continue
        # Sanitize filename & prevent directory traversal
        raw_name = os.path.basename(file.filename.replace("\\", "/"))
        safe_name = re.sub(r"[^a-zA-Z0-9_.-]", "_", raw_name)
        if not safe_name.strip("._"):
            safe_name = f"upload_{idx}"

        # Deduplicate filename if multiple files have identical names
        base_stem, ext = os.path.splitext(safe_name)
        ext = ext.lower()
        if ext == ".zip":
            raise IntakeError(
                f"File '{safe_name}': Direct ZIP uploads are not supported. Please extract files before uploading."
            )
        if ext not in ALLOWED_EXTENSIONS:
            raise IntakeError(
                f"Unsupported file type: '{ext}'. Allowed: {', '.join(sorted(ALLOWED_EXTENSIONS))}"
            )

        counter = 1
        candidate_name = safe_name
        while candidate_name in used_names:
            candidate_name = f"{base_stem}_{counter}{ext}"
            counter += 1
        safe_name = candidate_name
        used_names.add(safe_name)

        content = await file.read()
        file_size = len(content)
        total_bytes += file_size

        if file_size == 0:
            raise IntakeError(f"File '{safe_name}' is empty (0 bytes).")
        if file_size > MAX_FILE_BYTES:
            raise IntakeError(
                f"File '{safe_name}' exceeds 25 MB limit ({file_size / (1024 * 1024):.1f} MB)."
            )
        if total_bytes > MAX_TOTAL_BYTES:
            raise IntakeError(
                f"Total upload size exceeds 50 MB limit ({total_bytes / (1024 * 1024):.1f} MB)."
            )

        # Magic byte header validation to prevent extension spoofing
        if ext == ".pdf" and not content.startswith(b"%PDF-"):
            raise IntakeError(f"File '{safe_name}' does not match valid PDF header format.")
        elif ext in (".docx", ".xlsx") and not content.startswith(b"PK\x03\x04"):
            raise IntakeError(f"File '{safe_name}' is not a valid Office open XML document.")

        dest_path = raw_dir / safe_name
        with open(dest_path, "wb") as f:
            f.write(content)
        saved_file_paths.append(str(dest_path))

        # Extract text/units with defensive parsing
        file_manifest, extracted_lines = _extract_file(dest_path, safe_name, ext, content)
        manifest_files.append(file_manifest)
        all_extracted_lines.extend(extracted_lines)

        # Save extracted lines as json on disk
        with open(extracted_dir / f"{safe_name}.json", "w", encoding="utf-8") as f:
            json.dump(extracted_lines, f)

    # Determine overall unitKind and totalUnits
    total_units = sum(fm.units for fm in manifest_files)
    unit_kind: Literal["numbers", "rows", "lines", "tokens", "pages", "images", "unknown"]
    if not manifest_files:
        unit_kind = "lines"
    elif all(fm.unitKind == "numbers" for fm in manifest_files):
        unit_kind = "numbers"
    elif all(fm.unitKind == "rows" for fm in manifest_files):
        unit_kind = "rows"
    elif all(fm.unitKind == "pages" for fm in manifest_files):
        unit_kind = "pages"
    elif all(fm.unitKind == "images" for fm in manifest_files):
        unit_kind = "images"
    else:
        unit_kind = "lines"

    summary_str = (
        f"{len(manifest_files)} file(s), ~{total_units:,} {unit_kind}, {total_bytes / 1024:.1f} KB"
    )
    manifest = Manifest(
        files=manifest_files,
        totalUnits=total_units,
        unitKind=unit_kind,
        summary=summary_str,
    )

    # Save manifest.json
    with open(job_upload_dir / "manifest.json", "w", encoding="utf-8") as f:
        f.write(manifest.model_dump_json(indent=2))

    return manifest, saved_file_paths


def _extract_file(
    file_path: Path, filename: str, ext: str, content: bytes
) -> tuple[FileManifest, list[str]]:
    extracted_lines: list[str] = []
    unit_kind: str = "lines"
    preview = ""

    if ext in (".txt", ".md"):
        try:
            text = content.decode("utf-8")
        except UnicodeDecodeError:
            try:
                text = content.decode("latin-1")
            except Exception:
                text = content.decode("utf-8", errors="replace")
        extracted_lines = [line.strip() for line in text.splitlines() if line.strip()]
        if _is_number_list(extracted_lines):
            unit_kind = "numbers"
            units = sum(len(line.replace(",", " ").split()) for line in extracted_lines)
        else:
            unit_kind = "lines"
            units = len(extracted_lines)
        preview = text[:1200]

    elif ext == ".json":
        try:
            text = content.decode("utf-8", errors="replace")
            parsed = json.loads(text)
            if isinstance(parsed, list):
                extracted_lines = [json.dumps(item) for item in parsed]
                unit_kind = "rows"
                units = len(extracted_lines)
            else:
                extracted_lines = [text]
                unit_kind = "lines"
                units = 1
            preview = text[:1200]
        except Exception as e:
            raise IntakeError(f"File '{filename}' contains invalid or corrupted JSON: {e}")

    elif ext == ".csv":
        try:
            try:
                text = content.decode("utf-8")
            except UnicodeDecodeError:
                text = content.decode("latin-1")
            lines = text.splitlines()
            extracted_lines = lines
            unit_kind = "rows"
            units = max(0, len(lines) - 1) if len(lines) > 1 else len(lines)
            preview = "\n".join(lines[:8])
        except Exception as e:
            raise IntakeError(f"File '{filename}' cannot be parsed as CSV: {e}")

    elif ext == ".xlsx":
        try:
            wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True)
            sheet = wb.active
            rows: list[str] = []
            if sheet:
                for row in sheet.iter_rows(values_only=True):
                    rows.append(",".join("" if v is None else str(v) for v in row))
            extracted_lines = rows
            unit_kind = "rows"
            units = max(0, len(rows) - 1)
            preview = "\n".join(rows[:8])
        except Exception as e:
            raise IntakeError(f"File '{filename}' is corrupted or unreadable as XLSX: {e}")

    elif ext == ".pdf":
        try:
            reader = pypdf.PdfReader(io.BytesIO(content))
            if getattr(reader, "is_encrypted", False):
                raise IntakeError(
                    f"File '{filename}' is password-protected. Please upload an unprotected PDF."
                )
            units = len(reader.pages)
            unit_kind = "pages"
            pages_text = []
            for i, page in enumerate(reader.pages):
                p_text = page.extract_text() or ""
                pages_text.append(f"--- Page {i + 1} ---\n{p_text}")
                extracted_lines.append(p_text)
            if not any(t.strip() for t in extracted_lines):
                extracted_lines = ["(Scanned document without text layer)"]
                preview = "(Scanned PDF: text extraction requires OCR)"
            else:
                preview = "\n".join(pages_text[:3])[:1200]
        except IntakeError:
            raise
        except Exception as e:
            raise IntakeError(f"File '{filename}' is corrupted or unreadable as PDF: {e}")

    elif ext == ".docx":
        try:
            doc = docx.Document(io.BytesIO(content))
            extracted_lines = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
            units = len(extracted_lines)
            unit_kind = "lines"
            preview = "\n".join(extracted_lines[:10])[:1200]
        except Exception as e:
            raise IntakeError(f"File '{filename}' is corrupted or unreadable as DOCX: {e}")

    elif ext in (".png", ".jpg", ".jpeg", ".webp"):
        try:
            img = Image.open(io.BytesIO(content))
            img.verify()
            img = Image.open(io.BytesIO(content))
            unit_kind = "images"
            units = 1
            extracted_lines = [f"image:{filename}:{img.format}:{img.size[0]}x{img.size[1]}"]
            preview = f"Image {img.format} {img.size[0]}x{img.size[1]} px"
        except Exception as e:
            raise IntakeError(f"File '{filename}' is not a valid or readable image: {e}")

    else:
        text = content.decode("utf-8", errors="replace")
        extracted_lines = text.splitlines()
        units = len(extracted_lines)
        unit_kind = "lines"
        preview = text[:500]

    fm = FileManifest(
        name=filename,
        type=ext.lstrip("."),
        bytes=len(content),
        units=max(1, units),
        unitKind=unit_kind,  # type: ignore[arg-type]
        preview=preview,
    )
    return fm, extracted_lines


def _is_number_list(lines: list[str]) -> bool:
    if not lines:
        return False
    sample = lines[:30]
    total_tokens = 0
    number_tokens = 0
    for line in sample:
        tokens = line.replace(",", " ").split()
        for tok in tokens:
            total_tokens += 1
            if re.match(r"^[-+]?\d+(\.\d+)?$", tok.strip()):
                number_tokens += 1
    return total_tokens > 0 and (number_tokens / total_tokens) >= 0.8
